"""Write schedule service tabular outputs."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_mapping_outputs(output_dir: Path, output_stem: str, rows: list[dict[str, Any]]) -> tuple[Path, Path]:
    return write_table_outputs(output_dir, output_stem, rows, "Schedule Mapping")


def write_schedule_outputs(output_dir: Path, output_stem: str, rows: list[dict[str, Any]]) -> tuple[Path, Path]:
    return write_table_outputs(output_dir, output_stem, rows, "Generated Schedule")


def write_table_outputs(
    output_dir: Path,
    output_stem: str,
    rows: list[dict[str, Any]],
    sheet_title: str,
) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / f"{output_stem}.json"
    xlsx_path = output_dir / f"{output_stem}.xlsx"
    json_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_xlsx(xlsx_path, rows, sheet_title)
    return xlsx_path, json_path


def _write_xlsx(path: Path, rows: list[dict[str, Any]], sheet_title: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title
    headers = list(rows[0].keys()) if rows else []
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
        worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 70)

    workbook.save(path)
