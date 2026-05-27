"""Clean CCPP guide schedule workbooks into structured schedule data."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_INPUT = Path("04_data/schedule_sources/raw/ccpp guide schedule_260527.xlsx")
DEFAULT_OUTPUT_DIR = Path("04_data/schedule_sources/processed")
DEFAULT_OUTPUT_STEM = "ccpp_guide_schedule_260527_clean"
PREFERRED_SHEET = "Sheet1 (2)"


@dataclass(frozen=True)
class CleanScheduleRow:
    source_file: str
    source_sheet: str
    source_row: int
    row_type: str
    wbs_level: str
    wbs_path: str
    wbs_level_1: str
    wbs_level_2: str
    wbs_level_3: str
    wbs_level_4: str
    wbs_level_5: str
    activity_id: str
    activity_name: str
    activity_name_clean: str
    start_date: str
    finish_date: str
    search_text: str
    raw_wbs_level: str
    raw_trimmed_id: str
    raw_activity_name: str
    raw_start: str
    raw_finish: str
    raw_wbs_level_extra: str


def clean_guide_schedule(
    input_path: Path = DEFAULT_INPUT,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    output_stem: str = DEFAULT_OUTPUT_STEM,
) -> tuple[Path, Path]:
    input_path = input_path.resolve()
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(input_path, data_only=False, read_only=True)
    sheet_name = PREFERRED_SHEET if PREFERRED_SHEET in workbook.sheetnames else workbook.sheetnames[0]
    worksheet = workbook[sheet_name]

    rows = _clean_rows(input_path, worksheet)
    metadata = _build_metadata(input_path, workbook.sheetnames, sheet_name, worksheet.max_row, rows)

    json_path = output_dir / f"{output_stem}.json"
    xlsx_path = output_dir / f"{output_stem}.xlsx"
    _write_json(json_path, metadata, rows)
    _write_xlsx(xlsx_path, metadata, rows)
    return xlsx_path, json_path


def _clean_rows(input_path: Path, worksheet: Any) -> list[CleanScheduleRow]:
    current_wbs: dict[int, str] = {}
    clean_rows: list[CleanScheduleRow] = []

    for source_row, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        raw_wbs_level = _cell_to_text(values[0] if len(values) > 0 else None)
        raw_trimmed_id = _cell_to_text(values[1] if len(values) > 1 else None)
        raw_activity_name = _cell_to_text(values[2] if len(values) > 2 else None)
        raw_start = _cell_to_text(values[3] if len(values) > 3 else None)
        raw_finish = _cell_to_text(values[4] if len(values) > 4 else None)
        raw_wbs_level_extra = _cell_to_text(values[5] if len(values) > 5 else None)

        if not any([raw_wbs_level, raw_trimmed_id, raw_activity_name, raw_start, raw_finish, raw_wbs_level_extra]):
            continue

        row_type = _row_type(raw_wbs_level)
        wbs_level = raw_wbs_level if row_type == "wbs" else ""

        if row_type == "wbs":
            level = int(float(raw_wbs_level))
            current_wbs[level] = raw_trimmed_id
            for stale_level in [key for key in current_wbs if key > level]:
                del current_wbs[stale_level]

        wbs_path_parts = [current_wbs[key] for key in sorted(current_wbs) if current_wbs[key]]
        wbs_path = " > ".join(wbs_path_parts)
        activity_name_clean = _clean_activity_name(raw_activity_name)
        start_date = _parse_date(raw_start)
        finish_date = _parse_date(raw_finish)
        activity_id = raw_trimmed_id if row_type == "activity" else ""
        activity_name = raw_activity_name if row_type == "activity" else ""
        search_text = _make_search_text(wbs_path_parts, activity_id, activity_name_clean)

        clean_rows.append(
            CleanScheduleRow(
                source_file=input_path.name,
                source_sheet=worksheet.title,
                source_row=source_row,
                row_type=row_type,
                wbs_level=wbs_level,
                wbs_path=wbs_path,
                wbs_level_1=current_wbs.get(1, ""),
                wbs_level_2=current_wbs.get(2, ""),
                wbs_level_3=current_wbs.get(3, ""),
                wbs_level_4=current_wbs.get(4, ""),
                wbs_level_5=current_wbs.get(5, ""),
                activity_id=activity_id,
                activity_name=activity_name,
                activity_name_clean=activity_name_clean,
                start_date=start_date,
                finish_date=finish_date,
                search_text=search_text,
                raw_wbs_level=raw_wbs_level,
                raw_trimmed_id=raw_trimmed_id,
                raw_activity_name=raw_activity_name,
                raw_start=raw_start,
                raw_finish=raw_finish,
                raw_wbs_level_extra=raw_wbs_level_extra,
            )
        )

    return clean_rows


def _row_type(raw_wbs_level: str) -> str:
    if raw_wbs_level.strip().lower() == "activity":
        return "activity"
    if _is_numeric_level(raw_wbs_level):
        return "wbs"
    return "unknown"


def _is_numeric_level(value: str) -> bool:
    try:
        numeric = float(value)
    except ValueError:
        return False
    return numeric.is_integer()


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_date(value: str) -> str:
    cleaned = value.strip().rstrip("*").strip()
    if not cleaned:
        return ""

    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%b-%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat()
        except ValueError:
            continue
    return cleaned


def _clean_activity_name(value: str) -> str:
    cleaned = value.strip()
    cleaned = re.sub(r"\([^()]*NTP[^()]*\)", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned


def _make_search_text(wbs_path_parts: list[str], activity_id: str, activity_name: str) -> str:
    text = " ".join([*wbs_path_parts, activity_id, activity_name])
    return re.sub(r"\s+", " ", text).strip()


def _build_metadata(
    input_path: Path,
    sheet_names: list[str],
    selected_sheet: str,
    source_row_count: int,
    rows: list[CleanScheduleRow],
) -> dict[str, Any]:
    row_type_counts: dict[str, int] = {}
    for row in rows:
        row_type_counts[row.row_type] = row_type_counts.get(row.row_type, 0) + 1

    return {
        "source_file": str(input_path),
        "source_sheet_names": sheet_names,
        "selected_sheet": selected_sheet,
        "source_row_count": source_row_count,
        "clean_row_count": len(rows),
        "row_type_counts": row_type_counts,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "notes": [
            "Raw source columns are preserved as raw_* fields.",
            "start_date and finish_date are normalized when a known date format is detected.",
            "wbs_path is derived from the latest parent WBS rows above each row.",
        ],
    }


def _write_json(path: Path, metadata: dict[str, Any], rows: list[CleanScheduleRow]) -> None:
    payload = {
        "metadata": metadata,
        "rows": [asdict(row) for row in rows],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_xlsx(path: Path, metadata: dict[str, Any], rows: list[CleanScheduleRow]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    clean_sheet = workbook.create_sheet("Clean Schedule")
    activity_sheet = workbook.create_sheet("Activities Only")

    _write_summary_sheet(summary, metadata)
    _write_rows_sheet(clean_sheet, rows)
    _write_rows_sheet(activity_sheet, [row for row in rows if row.row_type == "activity"])
    workbook.save(path)


def _write_summary_sheet(worksheet: Any, metadata: dict[str, Any]) -> None:
    worksheet.append(["Field", "Value"])
    worksheet.append(["Source file", metadata["source_file"]])
    worksheet.append(["Selected sheet", metadata["selected_sheet"]])
    worksheet.append(["Source row count", metadata["source_row_count"]])
    worksheet.append(["Clean row count", metadata["clean_row_count"]])
    for row_type, count in metadata["row_type_counts"].items():
        worksheet.append([f"{row_type} rows", count])
    worksheet.append([])
    worksheet.append(["Notes", ""])
    for note in metadata["notes"]:
        worksheet.append(["", note])
    _format_sheet(worksheet)


def _write_rows_sheet(worksheet: Any, rows: list[CleanScheduleRow]) -> None:
    headers = list(CleanScheduleRow.__dataclass_fields__)
    worksheet.append(headers)
    for row in rows:
        worksheet.append([getattr(row, header) for header in headers])
    _format_sheet(worksheet)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _format_sheet(worksheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
        worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 60)


def main() -> None:
    parser = argparse.ArgumentParser(description="Clean a CCPP guide schedule workbook.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--output-stem", default=DEFAULT_OUTPUT_STEM)
    args = parser.parse_args()

    xlsx_path, json_path = clean_guide_schedule(args.input, args.output_dir, args.output_stem)
    print(f"Wrote cleaned workbook: {xlsx_path}")
    print(f"Wrote cleaned JSON: {json_path}")


if __name__ == "__main__":
    main()
