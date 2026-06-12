"""Load MDL source titles and classified CSV records."""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Any

import openpyxl

from common.text_normalizer import expand_abbreviation_terms, join_unique_texts, normalize_space
from mdl_service.models import DocumentTitle

TITLE_HEADERS = {"TITLE", "DOCUMENT DESCRIPTION", "DOCUMENT TITLE"}
DOCUMENT_NO_HEADERS = {"DOCUMENT NUMBER", "DOCUMENT NO", "DOCUMENT NO.", "DOC NO"}
EMBEDDING_FIELDS = (
    ("Title", "title"),
    ("Equipment", "equipment"),
    ("Building", "building"),
    ("System", "system"),
    ("Study/Survey", "study_survey"),
    ("Others", "others"),
    ("Deliverable", "deliverable"),
)


def list_excel_files(data_dir: str | Path, requested_file: str | None = None) -> list[Path]:
    """Return requested or discoverable MDL Excel files."""
    data_path = Path(data_dir)
    if requested_file:
        return [data_path / requested_file]
    return sorted(
        path
        for path in data_path.iterdir()
        if path.suffix.lower() in {".xlsx", ".xlsm"}
        and not path.name.startswith(("~$", "abbreviation"))
    )


def extract_titles_from_excel(filepath: str | Path) -> list[DocumentTitle]:
    """Extract MDL document titles from every worksheet with a recognized title header."""
    path = Path(filepath)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    results = []
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            title_col, document_no_col, header_row = _find_columns(rows[:10])
            if title_col is None or header_row is None:
                continue
            for row in rows[header_row + 1 :]:
                title = _cell_value(row, title_col)
                if not title or title.upper() in TITLE_HEADERS:
                    continue
                results.append(
                    DocumentTitle(
                        source_file=path.name,
                        sheet=sheet_name,
                        document_no=_cell_value(row, document_no_col),
                        title=title,
                    )
                )
    finally:
        workbook.close()
    return results


def load_ingest_records(csv_path: str | Path) -> list[dict[str, Any]]:
    """Load one classified MDL CSV into deduplicated normalized Neo4j records."""
    path = Path(csv_path)
    rows = _read_csv_rows(path)
    grouped_records: dict[str, dict[str, Any]] = {}
    title_order: list[str] = []
    for index, row in enumerate(rows):
        title = _clean(row.get("Title"))
        if not title:
            continue
        title_key = normalize_space(title).casefold()
        document_no = _clean(row.get("Document No")) or f"DOC_{index}"
        source_file = _clean(row.get("Source File"))
        record = {
            "project_names": _infer_project_name(source_file),
            "source_file": source_file,
            "document_no": document_no,
            "title": title,
            "equipment": _clean(row.get("Equipment")),
            "building": _clean(row.get("Building")),
            "system": _clean(row.get("System")),
            "study_survey": _clean(row.get("Study/Survey")),
            "others": _clean(row.get("Others")),
            "deliverable": _clean(row.get("Deliverable")),
        }
        if title_key not in grouped_records:
            grouped_records[title_key] = record
            title_order.append(title_key)
            continue
        grouped_records[title_key] = _merge_title_duplicate(grouped_records[title_key], record)
    records = []
    for title_key in title_order:
        record = grouped_records[title_key]
        record["doc_id"] = _build_title_doc_id(record["title"])
        record["text_content"] = build_embedding_text(record)
        records.append(record)
    return records


def dedupe_ingest_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Deduplicate ingest records globally by normalized title."""
    grouped_records: dict[str, dict[str, Any]] = {}
    title_order: list[str] = []
    for record in records:
        title = _clean(record.get("title"))
        if not title:
            continue
        title_key = normalize_space(title).casefold()
        normalized_record = {**record, "doc_id": _build_title_doc_id(title)}
        if title_key not in grouped_records:
            grouped_records[title_key] = normalized_record
            title_order.append(title_key)
            continue
        grouped_records[title_key] = _merge_title_duplicate(grouped_records[title_key], normalized_record)
    deduped_records = []
    for title_key in title_order:
        record = grouped_records[title_key]
        record["text_content"] = build_embedding_text(record)
        deduped_records.append(record)
    return deduped_records


def build_embedding_text(record: dict[str, Any]) -> str:
    """Build the MDL representation stored and embedded in Neo4j."""
    labeled_text = " | ".join(
        f"{label}: {value}"
        for label, field in EMBEDDING_FIELDS
        if (value := _clean(record.get(field)))
    )
    semantic_terms = [
        term
        for term in [
            _clean(record.get("title")),
            _clean(record.get("equipment")),
            _clean(record.get("building")),
            _clean(record.get("system")),
            _clean(record.get("study_survey")),
            _clean(record.get("others")),
            _clean(record.get("deliverable")),
        ]
        if term
    ]
    expanded_terms = _expanded_only_terms(semantic_terms)
    if expanded_terms:
        return f"{labeled_text} | Expanded Terms: {join_unique_texts(expanded_terms)}"
    return labeled_text


def _expanded_only_terms(terms: list[str]) -> list[str]:
    normalized_terms = _split_semantic_terms(terms)
    original_keys = {normalize_space(term).casefold() for term in normalized_terms}
    return [
        term
        for term in expand_abbreviation_terms(normalized_terms)
        if normalize_space(term).casefold() not in original_keys
    ]


def _split_semantic_terms(terms: list[str]) -> list[str]:
    split_terms: list[str] = []
    for term in terms:
        parts = [normalize_space(part) for part in str(term).split("|")]
        split_terms.extend(part for part in parts if part)
    return split_terms


def _merge_title_duplicate(existing: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    merged = dict(existing)
    for field in ("project_names", "source_file", "document_no"):
        merged[field] = _merge_unique_values(existing.get(field), incoming.get(field), separator="|")
    for field in ("equipment", "building", "system", "study_survey", "others", "deliverable"):
        merged[field] = _merge_unique_values(existing.get(field), incoming.get(field), separator=";")
    return merged


def _merge_unique_values(*values: Any, separator: str) -> str:
    unique_values: list[str] = []
    seen: set[str] = set()
    for raw in values:
        for part in str(raw or "").split(separator):
            cleaned = _clean(part)
            key = normalize_space(cleaned).casefold()
            if cleaned and key not in seen:
                unique_values.append(cleaned)
                seen.add(key)
    return f" {separator} ".join(unique_values)


def _build_title_doc_id(title: str) -> str:
    normalized_title = normalize_space(title).casefold()
    digest = hashlib.sha1(normalized_title.encode("utf-8")).hexdigest()[:12]
    return f"title_{digest}"


def _infer_project_name(source_file: str) -> str:
    stem = Path(source_file).stem if source_file else ""
    return stem[:-4] if stem.endswith("_MDL") else stem


def _find_columns(rows: list[tuple[Any, ...]]) -> tuple[int | None, int | None, int | None]:
    for row_index, row in enumerate(rows):
        title_col = None
        document_no_col = None
        for column_index, cell in enumerate(row):
            value = _clean(cell).upper()
            if value in TITLE_HEADERS:
                title_col = column_index
            if value in DOCUMENT_NO_HEADERS:
                document_no_col = column_index
        if title_col is not None:
            return title_col, document_no_col, row_index
    return None, None, None


def _cell_value(row: tuple[Any, ...], index: int | None) -> str:
    return "" if index is None or index >= len(row) else _clean(row[index])


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            with open(path, newline="", encoding=encoding) as file:
                return list(csv.DictReader(file))
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("utf-8-sig", b"", 0, 1, f"Unable to decode {path}")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text
